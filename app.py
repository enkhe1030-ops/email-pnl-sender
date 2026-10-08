import re
import smtplib
import time
import uuid
import random
import io
import os
import json
import hashlib
from datetime import datetime
from zoneinfo import ZoneInfo  
from email.message import EmailMessage
from email.utils import formatdate, formataddr
import pandas as pd
import streamlit as st
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ============================================================
# AUTHENTICATION, CONFIG & LOGGING SYSTEM
# ============================================================

USERS_FILE = "users.json"
LOGS_FILE = "logs.csv"
CREDENTIALS_FILE = "saved_credentials.json"

def hash_password(password):
    return hashlib.sha256(password.encode()).hexdigest()

def load_users():
    default_users = {
        "admin": {
            "password_hash": hash_password("admin123"),
            "role": "admin"
        },
        "Anujin": {
            "password_hash": hash_password("Anuujin111"),
            "role": "user"
        },
        "Bayartsetseg": {
            "password_hash": hash_password("Bayartsetseg7854"),
            "role": "user"
        },
        "Baigalmaa": {
            "password_hash": hash_password("Baigalmaa"),
            "role": "user"
        },
        "Buyka": {
            "password_hash": hash_password("Buyka1234"),
            "role": "user"
        },
        "Jagana": {
            "password_hash": hash_password("Jagana333"),
            "role": "user"
        },
        "Namuun": {
            "password_hash": hash_password("Namuun999"),
            "role": "user"
        },
        "Narangarav": {
            "password_hash": hash_password("Narangarav987"),
            "role": "user"
        },
        "Nyamtsetseg": {
            "password_hash": hash_password("Nyamtsetseg2247"),
            "role": "user"
        },
        "Odsuren": {
            "password_hash": hash_password("Odsuren369"),
            "role": "user"
        },
        "otgonbileg": {
            "password_hash": hash_password("Otgonbileg4567"),
            "role": "user"
        },
        "Sunderiya": {
            "password_hash": hash_password("Sunderiya1234"),
            "role": "user"
        },
        "Selenge": {
            "password_hash": hash_password("Selenge4785"),
            "role": "user"
        },
        "Uugantsetseg": {
            "password_hash": hash_password("Uugantsetseg222"),
            "role": "user"
        },
        "Uyaraa": {
            "password_hash": hash_password("Uyaraa1265"),
            "role": "user"
        },
        "Khulan": {
            "password_hash": hash_password("Khulan6157"),
            "role": "user"
        },
        "Tsetsegmaa": {
            "password_hash": hash_password("Tsetsegmaa8541"),
            "role": "user"
        },
        "Chuluuntsetseg": {
            "password_hash": hash_password("Chuluuntsetseg976"),
            "role": "user"
        },
        "Erdnetsetseg": {
            "password_hash": hash_password("Erdnetsetseg7432"),
            "role": "user"
        },
        "Enkhtuul": {
            "password_hash": hash_password("Etu0610"),
            "role": "user"
        },
        "Jargal": {
            "password_hash": hash_password("Jargal2345"),
            "role": "user"
        },
        "Maralmaa": {
            "password_hash": hash_password("Maralmaa9711"),
            "role": "user"
        },
        "Enkhtuya": {
            "password_hash": hash_password("Enkhtuya2247"),
            "role": "user"
        },
        "Unurjargal": {
            "password_hash": hash_password("Unurjargal3346"),
            "role": "user"
        },
        "Enkhamgalan": {
            "password_hash": hash_password("Enkhamgalan7821"),
            "role": "user"
        },
        "Javzandulam": {
            "password_hash": hash_password("Javzandulam3954"),
            "role": "user"
        },
        "Enkhmaa": {
            "password_hash": hash_password("Enkhmaa3399"),
            "role": "user"
        },
        "Enkhtsetseg": {
            "password_hash": hash_password("Enkhtsetseg333999"),
            "role": "user"
        }
    }
    
    if not os.path.exists(USERS_FILE):
        save_users(default_users)
        return default_users
    
    try:
        with open(USERS_FILE, "r", encoding="utf-8") as f:
            users_dict = json.load(f)
            updated = False
            for k, v in default_users.items():
                if k not in users_dict:
                    users_dict[k] = v
                    updated = True
            if updated:
                save_users(users_dict)
            return users_dict
    except Exception:
        save_users(default_users)
        return default_users

def save_users(users_dict):
    with open(USERS_FILE, "w", encoding="utf-8") as f:
        json.dump(users_dict, f, ensure_ascii=False, indent=4)

def load_smtp_credentials():
    if os.path.exists(CREDENTIALS_FILE):
        try:
            with open(CREDENTIALS_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return {"sender_email": "no.reply.mongolian.airlines@gmail.com", "app_password": ""}
    return {"sender_email": "no.reply.mongolian.airlines@gmail.com", "app_password": ""}

def save_smtp_credentials(sender_email, app_password):
    data = {
        "sender_email": sender_email,
        "app_password": app_password
    }
    with open(CREDENTIALS_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=4)

def add_log(username, action, flight_no="", flight_date="", details=""):
    now_str = get_ubn_now()
    log_data = {
        "Timestamp": now_str,
        "Username": username,
        "Action": action,
        "Flight No": flight_no if flight_no else "-",
        "Flight Date": flight_date if flight_date else "-",
        "Details": details
    }
    df_new = pd.DataFrame([log_data])
    if not os.path.exists(LOGS_FILE):
        df_new.to_csv(LOGS_FILE, index=False, encoding="utf-8-sig")
    else:
        try:
            df_existing = pd.read_csv(LOGS_FILE, encoding="utf-8-sig", on_bad_lines="skip")
            if list(df_existing.columns) != list(df_new.columns):
                df_new.to_csv(LOGS_FILE, index=False, encoding="utf-8-sig")
            else:
                df_new.to_csv(LOGS_FILE, mode="a", header=False, index=False, encoding="utf-8-sig")
        except Exception:
            df_new.to_csv(LOGS_FILE, index=False, encoding="utf-8-sig")

# ============================================================
# HELPER FUNCTIONS
# ============================================================

def get_ubn_now():
    return datetime.now(ZoneInfo("Asia/Ulaanbaatar")).strftime("%Y-%m-%d %H:%M:%S")

def format_date_custom(raw_date_str):
    if not raw_date_str:
        return "", "", ""

    months = {
        "JAN": 1, "FEB": 2, "MAR": 3, "APR": 4, "MAY": 5, "JUN": 6,
        "JUL": 7, "AUG": 8, "SEP": 9, "OCT": 10, "NOV": 11, "DEC": 12
    }

    m = re.search(r"(\d{2})([A-Z]{3})(\d{2,4})?", raw_date_str.upper())
    if m:
        day = int(m.group(1))
        month_str = m.group(2)
        year_str = m.group(3) if m.group(3) else str(datetime.now(ZoneInfo("Asia/Ulaanbaatar")).year)

        if len(year_str) == 2:
            year_str = "20" + year_str

        year = int(year_str)
        month = months.get(month_str, 1)

        dt = datetime(year, month, day)
        mn_dot_date = dt.strftime("%Y.%m.%d")
        mn_dash_date = dt.strftime("%Y-%m-%d")
        en_date = dt.strftime("%B %d, %Y")

        return mn_dot_date, mn_dash_date, en_date

    return raw_date_str, raw_date_str, raw_date_str

def clean_and_decode_email(raw_str):
    raw_str = raw_str.strip()
    lang = "N/A"
    
    upper_str = raw_str.upper()
    if "/EN" in upper_str or "-EN" in upper_str:
        lang = "EN"
        raw_str = re.sub(r'[/_ \-]EN$', '', raw_str, flags=re.IGNORECASE)
    elif "/MN" in upper_str or "-MN" in upper_str:
        lang = "MN"
        raw_str = re.sub(r'[/_ \-]MN$', '', raw_str, flags=re.IGNORECASE)

    if "E+" in raw_str.upper():
        parts = re.split(r'E\+', raw_str, flags=re.IGNORECASE)
        if len(parts) > 1:
            raw_str = parts[1]

    email = raw_str.replace("//", "@").replace("..", "_").replace("./", "-")
    email = re.sub(r'^[^\w\.\-\+]+', '', email)
    email = re.sub(r'[^\w\.\-\+]+$', '', email)
    email = email.strip().lower()

    return email, lang

def is_valid_email(email):
    if not email or "@" not in email or " " in email:
        return False
    email_regex = r'^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$'
    return bool(re.match(email_regex, email))

# ============================================================
# PARSE AMADEUS PNL (Зөвхөн зорчигч унших зориулалттай)
# ============================================================

def parse_pnl(text):
    lines = text.splitlines()
    passenger_pattern = re.compile(
        r"^\s*(\d{1,3})[A-Z]?\s+(?:\*\d{1,2}|\d{1,2})?\s*(.+?)\s+([A-Z0-9]{5,8})(?:\s+([A-Z]))?\s+([A-Z]{2})\s+(\d{2}[A-Z]{3})\s*([A-Z0-9]+)?",
        re.IGNORECASE
    )
    ticket_pattern = re.compile(r"FA\s+PAX\s+(\d{3}-\d{9,10})", re.IGNORECASE)

    raw_passengers = []
    current = None

    for line in lines:
        passenger_match = passenger_pattern.search(line)
        if passenger_match:
            if current is not None:
                raw_passengers.append(current)

            pnl_no = passenger_match.group(1)
            pax_name_raw = passenger_match.group(2).strip()
            pax_name_cleaned = re.sub(r"^\s*\*?\d{1,2}\s*", "", pax_name_raw)

            pnr = passenger_match.group(3).upper()
            booking_class = (passenger_match.group(4) or "-").upper()
            status_code = (passenger_match.group(5) or "-").upper()
            booking_date = (passenger_match.group(6) or "-").upper()
            office_code = (passenger_match.group(7) or "-").upper()

            current = {
                "PNLNo": pnl_no,
                "Passenger Name": pax_name_cleaned,
                "PNR": pnr,
                "Class": booking_class,
                "Status Code": status_code,
                "Booking Date": booking_date,
                "OfficeCode": office_code,
                "TicketNo": "-",
                "emails_dict": {}
            }
            continue

        if current is not None:
            line_str = line.strip()
            tkt_match = ticket_pattern.search(line_str)
            if tkt_match:
                current["TicketNo"] = tkt_match.group(1)

            tokens = line_str.split()
            for token in tokens:
                if "//" in token or "@" in token or "E+" in token.upper():
                    email, lang = clean_and_decode_email(token)
                    if is_valid_email(email):
                        if email not in current["emails_dict"] or current["emails_dict"][email] == "N/A":
                            current["emails_dict"][email] = lang

    if current is not None:
        raw_passengers.append(current)

    formatted_records, missing_data_records = [], []
    valid_idx, missing_idx = 1, 1

    valid_statuses = ["HK", "TK", "SA", "RR", "UN"]

    for pax in raw_passengers:
        emails_list = list(pax["emails_dict"].keys())
        langs_list = list(pax["emails_dict"].values())

        if "MN" in langs_list:
            primary_lang = "MN"
        elif "EN" in langs_list:
            primary_lang = "EN"
        else:
            primary_lang = "N/A"

        status = pax["Status Code"]
        missing_reasons = []

        if pax["TicketNo"] == "-":
            missing_reasons.append("ТК Т дугааргүй")

        if status == "UC":
            missing_reasons.append("Статус UC")
        elif status not in valid_statuses:
            missing_reasons.append(f"{status if status != '-' else 'Нэг ч'} статусгүй/буруу статус")
        
        if not emails_list:
            missing_reasons.append("Имэйл хаяггүй")

        record = {
            "Selected": True,
            "SeqNo": "",
            "PNLNo": pax["PNLNo"],
            "Passenger Name": pax["Passenger Name"],
            "PNR": pax["PNR"],
            "Class": pax["Class"],
            "StatusCode": pax["Status Code"],
            "BookingDate": pax["Booking Date"],
            "OfficeCode": pax["OfficeCode"],
            "TicketNo": pax["TicketNo"],
            "Email": ", ".join(emails_list) if emails_list else "ОЛДООГҮЙ",
            "EmailList": emails_list,
            "Language": primary_lang,
            "SendStatus": "Not Processed",
            "SentTime": "-",
            "MissingReason": ", ".join(missing_reasons)
        }

        if missing_reasons:
            record["SeqNo"] = str(missing_idx)
            missing_data_records.append(record)
            missing_idx += 1
        else:
            record["SeqNo"] = str(valid_idx)
            formatted_records.append(record)
            valid_idx += 1

    return formatted_records, missing_data_records

# ============================================================
# TEMPLATE GENERATOR
# ============================================================

def generate_email_text_base(target_lang, flight_info):
    pax_name = "{PAX_NAME}"
    pnr_code = "{PNR}"
    tkt_no = "{TICKET_NO}"

    flt_no = flight_info['flight'] if flight_info['flight'] else ""
    flt_date = flight_info['date'] if flight_info['date'] else ""
    dep_time = flight_info['dep_time']
    arr_time = flight_info['arr_time']
    reason_mn = flight_info.get('reason_mn', '')
    reason_en = flight_info.get('reason_en', '')
    status_type = flight_info['status_type']

    mn_dot_date, mn_dash_date, en_date = format_date_custom(flt_date) if flt_date else ("", "", "")

    no_reply_footer_mn = "\n\n--------------------------------------------------\nЭнэхүү мэйл нь автоматаар илгээгдэж буй тул хариу бичих шаардлагагүй."
    no_reply_footer_en = "\n\n--------------------------------------------------\nThis is an automated message, please do not reply to this email."

    if target_lang == "MN":
        if status_type == "CANCEL":
            subject = f"{mn_dot_date} –ний {flt_no} нислэг цуцлагдсан тухай мэдэгдэл".strip()
            body = f"Хүндэт {pax_name},\n\nТаны {mn_dash_date}-ны өдрийн {flt_no} дугаартай нислэг цуцлагдсан болохыг үүгээр мэдэгдэж байна.\n\nЗОРЧИГЧИЙН МЭДЭЭЛЭЛ:\n• Зорчигчийн нэр: {pax_name}\n• Захиалгын дугаар (PNR): {pnr_code}\n• Тийзийн дугаар: {tkt_no}\n\nЦУЦЛАГДСАН НИСЛЭГИЙН МЭДЭЭЛЭЛ:\n• Нислэг: {flt_no}\n• Огноо: {flt_date}"
            if dep_time: body += f"\n• Нисэх цаг: {dep_time}"
            if arr_time: body += f"\n• Буух цаг: {arr_time}"
            if reason_mn: body += f"\n• Шалтгаан: {reason_mn}"
            body += f"\n\nТийз буцаалт болон өөр өдрийн нислэгээр тийзээ өөрчилж баталгаажуулах талаар тийз худалдан авсан аяллын агентлаг эсхүл тийз олгосон газартайгаа аль болох хурдан хугацаанд холбогдоно уу.\n\nДээрх өөрчлөлтөөс шалтгаалан Танд хүндрэл, чирэгдэл учруулж байгаад хүлцэл өчье.\n\nХүндэтгэсэн,\nМИАТ ТӨХК{no_reply_footer_mn}"
        else:
            subject = f"{mn_dot_date} –ний {flt_no} нислэгийн хуваарийн өөрчлөлтийн тухай мэдэгдэл".strip()
            body = f"Хүндэт {pax_name},\n\nТаны {mn_dash_date}-ны өдрийн {flt_no} дугаартай нислэгийн цагийн хуваарьт өөрчлөлт орсон болохыг үүгээр мэдэгдэж байна.\n\nЗОРЧИГЧИЙН МЭДЭЭЛЭЛ:\n• Зорчигчийн нэр: {pax_name}\n• Захиалгын дугаар (PNR): {pnr_code}\n• Тийзийн дугаар: {tkt_no}\n\nШИНЭ НИСЛЭГИЙН МЭДЭЭЛЭЛ:\n• Нислэг: {flt_no}\n• Огноо: {flt_date}"
            if dep_time: body += f"\n• Нисэх цаг: {dep_time}"
            if arr_time: body += f"\n• Буух цаг: {arr_time}"
            body += f"\n\nТаны тийзийн төлөв байдал болон шинэ нислэгийн мэдээллийг баталгаажуулахын тулд тийз худалдан авсан аяллын агентлаг эсхүл тийз олгосон газартайгаа аль болох хурдан хугацаанд холбогдоно уу.\n\nДээрх өөрчлөлтөөс шалтгаалан Танд хүндрэл, чирэгдэл учруулж байгаад хүлцэл өчье.\n\nХүндэтгэсэн,\nМИАТ ТӨХК{no_reply_footer_mn}"
    else: 
        if status_type == "CANCEL":
            subject = f"Flight Cancellation Notification - {flt_no} - {en_date}".strip()
            body = f"Dear {pax_name},\n\nWe regret to inform you that your flight {flt_no}, scheduled for {en_date}, has been cancelled.\n\nPASSENGER DETAILS:\n- Passenger Name: {pax_name}\n- Booking Reference (PNR): {pnr_code}\n- Ticket Number: {tkt_no}\n\nCANCELLED FLIGHT DETAILS:\n- Flight: {flt_no}\n- Date: {flt_date}"
            if dep_time: body += f"\n- Departure Time: {dep_time}"
            if arr_time: body += f"\n- Arrival Time: {arr_time}"
            if reason_en: body += f"\n- Reason: {reason_en}"
            body += f"\n\nFor ticket refund or to change and confirm your ticket for a flight on another date, please contact your travel agent or ticket issuing office as soon as possible.\n\nBest regards,\nMIAT Mongolian Airlines{no_reply_footer_en}"
        else:
            subject = f"Flight Schedule Change Notification - {flt_no} – {en_date}"
            body = f"Dear {pax_name},\n\nWe regret to inform you of a schedule change for your flight {flt_no} on {en_date}.\n\nPASSENGER DETAILS:\n- Passenger Name: {pax_name}\n- Booking Reference (PNR): {pnr_code}\n- Ticket Number: {tkt_no}\n\nNEW FLIGHT SCHEDULE DETAILS:\n- Flight: {flt_no}\n- Date: {flt_date}"
            if dep_time: body += f"\n- Departure Time: {dep_time}"
            if arr_time: body += f"\n- Arrival Time: {arr_time}"
            body += f"\n\nPlease contact your travel agent or issuing office as soon as possible to confirm your flight details.\n\nBest regards,\nMIAT Mongolian Airlines{no_reply_footer_en}"

    return subject, body

def text_to_html(plain_text):
    formatted = plain_text.replace('\n', '<br>')
    formatted = re.sub(r'(\b[A-Z-0-9А-ЯӨҮөү\s]+:)', r'<b>\1</b>', formatted)
    
    html_wrapper = f"""
    <!DOCTYPE html>
    <html lang="mn">
    <head>
        <meta charset="utf-8">
        <meta name="viewport" content="width=device-width, initial-scale=1.0">
        <style>
            body {{
                font-family: Arial, Helvetica, sans-serif;
                font-size: 14px;
                color: #222222;
                line-height: 1.6;
                background-color: #f9f9f9;
                margin: 0;
                padding: 20px;
            }}
            .container {{
                max-width: 600px;
                margin: 0 auto;
                padding: 25px;
                border: 1px solid #e0e0e0;
                border-radius: 8px;
                background-color: #ffffff;
                box-shadow: 0 2px 4px rgba(0,0,0,0.05);
            }}
        </style>
    </head>
    <body>
        <div class="container">
            {formatted}
        </div>
    </body>
    </html>
    """
    return html_wrapper

def render_custom_template(template_text, record):
    pax_name = record.get('Passenger Name', '{PAX_NAME}') if record else "{PAX_NAME}"
    pnr_code = record.get('PNR', '{PNR}') if record else "{PNR}"
    tkt_no = record.get('TicketNo', '{TICKET_NO}') if record else "{TICKET_NO}"

    rendered = template_text.replace("{PAX_NAME}", pax_name)
    rendered = rendered.replace("{PNR}", pnr_code)
    rendered = rendered.replace("{TICKET_NO}", tkt_no)
    return rendered

# ============================================================
# HIGH DELIVERABILITY SMTP SENDER
# ============================================================

def send_email_smtp(sender_email, app_password, recipient_email, recipient_name, subject, plain_text, html_text):
    msg = EmailMessage()
    msg['Date'] = formatdate(localtime=True)
    
    unique_id = uuid.uuid4().hex
    msg['Message-ID'] = f"<{unique_id}.notification@gmail.com>"
    
    msg['From'] = formataddr(("MIAT Mongolian Airlines Notification", sender_email))
    msg['To'] = formataddr((recipient_name, recipient_email))
    msg['Reply-To'] = "MIAT No-Reply <no-reply@miat.com>"
    
    msg['Subject'] = subject
    
    msg['User-Agent'] = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) MIAT-NotificationManager/2.5"
    msg['X-Mailer'] = "MIAT-NotificationManager/2.5"
    msg['Auto-Submitted'] = "auto-generated" 
    msg['Precedence'] = "bulk"
    msg['List-Unsubscribe'] = "<mailto:no-reply@miat.com?subject=unsubscribe>"

    msg.set_content(plain_text)
    msg.add_alternative(html_text, subtype='html')

    with smtplib.SMTP('smtp.gmail.com', 587) as server:
        server.ehlo()
        server.starttls()
        server.ehlo()
        server.login(sender_email, app_password)
        server.send_message(msg)

# ============================================================
# EXCEL GENERATOR
# ============================================================

def create_formatted_excel(records, missing_records):
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Passenger List"
    
    all_records = records + missing_records
    if all_records:
        df1 = pd.DataFrame(all_records)
        cols_to_drop = ["Selected", "SeqNo", "EmailList"]
        df1 = df1.drop(columns=[c for c in cols_to_drop if c in df1.columns])
    else:
        df1 = pd.DataFrame(columns=["PNLNo", "Passenger Name", "PNR", "Class", "StatusCode", "BookingDate", "OfficeCode", "TicketNo", "Email", "Language", "SendStatus", "SentTime", "MissingReason"])

    ws1.append(list(df1.columns))
    for row in df1.itertuples(index=False):
        ws1.append(list(row))
        
    ws2 = wb.create_sheet(title="PNL Summary")
    target_cols = [
        ("PNLNo", "PNL No"),
        ("Passenger Name", "Passenger Name"),
        ("PNR", "PNR"),
        ("Class", "Class"),
        ("BookingDate", "Booking date"),
        ("TicketNo", "Ticket NO"),
        ("OfficeCode", "Office code")
    ]
    ws2.append([col[1] for col in target_cols])
    
    def safe_pnl_sort_key(item):
        pnl_val = str(item.get("PNLNo", "0")).strip()
        digits = re.sub(r'\D', '', pnl_val)
        return int(digits) if digits else 999999

    sorted_pnl_records = sorted(all_records, key=safe_pnl_sort_key)
    for rec in sorted_pnl_records:
        row_data = [rec.get(col[0], "-") for col in target_cols]
        ws2.append(row_data)

    calibri_font = Font(name="Calibri", size=10, bold=False)
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
    )
    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')

    for ws in [ws1, ws2]:
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
        ws.page_margins.top = 0.5
        ws.page_margins.bottom = 0.5
        ws.page_margins.left = 0.25
        ws.page_margins.right = 0.25
        
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_center

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.font = calibri_font
                cell.border = thin_border
                cell.alignment = align_left

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    return excel_buffer

# ============================================================
# CLEAR ALL CALLBACK FUNCTION
# ============================================================

def clear_all_data():
    st.session_state.records = []
    st.session_state.missing_records = []
    st.session_state.pnl_text = ""
    st.session_state.pnl_textarea = ""
    st.session_state.custom_templates = {"MN": {"subject": "", "text": ""}, "EN": {"subject": "", "text": ""}}
    st.session_state.edit_mode = False

    st.session_state.input_flt_no = ""
    st.session_state.input_flt_date = ""
    st.session_state.input_dep_time = ""
    st.session_state.input_arr_time = ""
    st.session_state.input_reason_mn = ""
    st.session_state.input_reason_en = ""

# ============================================================
# STREAMLIT INITIALIZATION & AUTHENTICATION
# ============================================================

st.set_page_config(page_title="MIAT Flight Notification System", layout="wide")

if "logged_in" not in st.session_state:
    st.session_state.logged_in = False
if "user_info" not in st.session_state:
    st.session_state.user_info = None

# AUTHENTICATION SCREEN
if not st.session_state.logged_in:
    st.markdown("<h2 style='text-align: center;'>✈️ MIAT Flight Notification System</h2>", unsafe_allow_html=True)
    st.markdown("<h4 style='text-align: center;'>Системд нэвтрэх</h4>", unsafe_allow_html=True)
    
    col_l1, col_l2, col_l3 = st.columns([1, 1, 1])
    with col_l2:
        username_input = st.text_input("Нэвтрэх нэр (Username)")
        password_input = st.text_input("Нууц үг (Password)", type="password")
        
        if st.button("Нэвтрэх", type="primary", use_container_width=True):
            users = load_users()
            hashed = hash_password(password_input)
            
            if username_input in users and users[username_input]["password_hash"] == hashed:
                # Хэрэглэгч солигдоход эсвэл шинээр нэвтрэхэд өмнөх ажлын талбарын мэдээллийг бүрэн цэвэрлэх
                st.session_state.records = []
                st.session_state.missing_records = []
                st.session_state.pnl_text = ""
                st.session_state.pnl_textarea = ""
                st.session_state.custom_templates = {"MN": {"subject": "", "text": ""}, "EN": {"subject": "", "text": ""}}
                st.session_state.edit_mode = False
                st.session_state.input_flt_no = ""
                st.session_state.input_flt_date = ""
                st.session_state.input_dep_time = ""
                st.session_state.input_arr_time = ""
                st.session_state.input_reason_mn = ""
                st.session_state.input_reason_en = ""

                st.session_state.logged_in = True
                st.session_state.user_info = {
                    "username": username_input,
                    "role": users[username_input]["role"]
                }
                add_log(username_input, "Нэвтэрсэн", details="Амжилттай нэвтэрлээ")
                st.success("Амжилттай нэвтэрлээ!")
                st.rerun()
            else:
                st.error("Нэвтрэх нэр эсвэл нууц үг буруу байна!")
    st.stop()

# ============================================================
# MAIN APPLICATION (LOGGED IN)
# ============================================================

if "records" not in st.session_state:
    st.session_state.records = []
if "missing_records" not in st.session_state:
    st.session_state.missing_records = []
if "pnl_text" not in st.session_state:
    st.session_state.pnl_text = ""
if "pnl_textarea" not in st.session_state:
    st.session_state.pnl_textarea = ""
if "custom_templates" not in st.session_state:
    st.session_state.custom_templates = {"MN": {"subject": "", "text": ""}, "EN": {"subject": "", "text": ""}}
if "edit_mode" not in st.session_state:
    st.session_state.edit_mode = False

if "input_flt_no" not in st.session_state:
    st.session_state.input_flt_no = ""
if "input_flt_date" not in st.session_state:
    st.session_state.input_flt_date = ""
if "input_dep_time" not in st.session_state:
    st.session_state.input_dep_time = ""
if "input_arr_time" not in st.session_state:
    st.session_state.input_arr_time = ""
if "input_reason_mn" not in st.session_state:
    st.session_state.input_reason_mn = ""
if "input_reason_en" not in st.session_state:
    st.session_state.input_reason_en = ""

# --- SIDEBAR: AUTHENTICATION & SAVED CREDENTIALS ---
with st.sidebar:
    st.markdown(f"### 👤 Хэрэглэгч: **{st.session_state.user_info['username']}**")
    st.caption(f"Эрх: {st.session_state.user_info['role'].upper()}")
    
    if st.button("🚪 Системээс гарах"):
        add_log(st.session_state.user_info['username'], "Гарсан", details="Системээс гарлаа")
        st.session_state.logged_in = False
        st.session_state.user_info = None
        # Системээс гарах үед бас ажлын талбарын утгуудыг цэвэрлэх
        st.session_state.records = []
        st.session_state.missing_records = []
        st.session_state.pnl_textarea = ""
        st.rerun()

    st.divider()
    st.header("🔑 Илгээгчийн Тохиргоо")
    
    saved_creds = load_smtp_credentials()
    
    if st.session_state.user_info["role"] == "admin":
        st.caption("⚙️ Админ тохиргоо (Санах ойд хадгалагдана)")
        sender_email = st.text_input("Илгээх Gmail Хаяг", value=saved_creds.get("sender_email", ""))
        app_password = st.text_input("Gmail App Password", value=saved_creds.get("app_password", ""), type="password", help="16 оронтой App Password оруулаад 'Хадгалах' товч дарна.")
        
        if st.button("💾 Хадгалах"):
            save_smtp_credentials(sender_email, app_password)
            st.success("App Password амжилттай хадгалагдлаа!")
            st.rerun()
    else:
        sender_email = saved_creds.get("sender_email", "")
        app_password = saved_creds.get("app_password", "")
        st.info(f"📧 **Илгээгч:** {sender_email}")
        if app_password:
            st.success("✅ App Password бэлэн хадгалагдсан байна.")
        else:
            st.warning("⚠️ АДМИН App Password хадгалаагүй байна.")

    st.divider()
    st.markdown("### 🛡 Inbox-д оруулах хамгаалалт:")
    st.caption("1. **No-Reply тохиргоо**: Зорчигч таны хувийн Gmail рүү хариу мэйл бичих боломжгүй.")
    st.caption("2. **Blacklist-ээс хамгаалах**: Код нь мэйл хооронд санамсаргүй хугацааны хүлээлт (2-4.5сек) авч илгээнэ.")
    st.caption("3. **Лимит**: Энгийн Gmail өдөрт 500 хүртэл мэйл илгээх лимиттэйг анхаарна уу.")

st.title("✈️ MIAT Flight Notification System (Inbox & Anti-Spam Safe)")

# --- MAIN LAYOUT ---
col1, col2 = st.columns([1, 1])

with col1:
    st.subheader("1. Amadeus Passenger Name List (PNL)")
    pnl_input = st.text_area("PNL Эх текст хуулах:", height=250, key="pnl_textarea")

    col_btn1, col_btn2 = st.columns([1, 1])
    if col_btn1.button("Extract PNL", type="primary"):
        if pnl_input.strip():
            st.session_state.pnl_text = pnl_input
            st.session_state.records, st.session_state.missing_records = parse_pnl(pnl_input)
            st.success("PNL амжилттай уншигдлаа!")
            st.rerun()
        else:
            st.warning("PNL текстээ оруулна уу.")
            
    col_btn2.button("Clear All", on_click=clear_all_data)

with col2:
    st.subheader("2. Нислэгийн Мэдээлэл")
    
    status_type = st.radio("Мэдэгдлийн төрөл:", ["CHANGE", "CANCEL"], format_func=lambda x: "Schedule Change (Өөрчлөгдсөн)" if x == "CHANGE" else "Flight Cancelled (Цуцлагдсан)", horizontal=True)
    
    c1, c2 = st.columns(2)
    flt_no = c1.text_input("Flight No:", key="input_flt_no", placeholder="OM297")
    flt_date = c2.text_input("Date:", key="input_flt_date", placeholder="01NOV")
    
    c4, c5 = st.columns(2)
    dep_time = c4.text_input("Dep Time:", placeholder="10:00", key="input_dep_time")
    arr_time = c5.text_input("Arr Time:", placeholder="14:30", key="input_arr_time")
    
    reason_mn = ""
    reason_en = ""
    if status_type == "CANCEL":
        rc1, rc2 = st.columns(2)
        reason_mn = rc1.text_input("Шалтгаан (MN):", placeholder="Техникийн саатал...", key="input_reason_mn")
        reason_en = rc2.text_input("Reason (EN):", placeholder="Technical reason...", key="input_reason_en")

flight_info = {
    "flight": flt_no,
    "date": flt_date,
    "dep_time": dep_time,
    "arr_time": arr_time,
    "reason_mn": reason_mn,
    "reason_en": reason_en,
    "status_type": status_type
}

st.divider()

# --- PASSENGER TABLES, PREVIEW, LOGS & USER MANAGEMENT ---
tab1, tab2, tab3, tab4, tab5 = st.tabs(["📋 Идэвхтэй Зорчигчид", "⚠️ Мэдээлэл Дутуу Зорчигчид", "👁️ Email Preview", "📜 Үйл ажиллагааны түүх", "👤 Хэрэглэгчдийн Удирдлага"])

with tab1:
    if st.session_state.records:
        col_s1, col_s2, col_s3 = st.columns([1, 1, 3])
        if col_s1.button("☑️ Бүгдийг сонгох"):
            for r in st.session_state.records:
                r["Selected"] = True
            st.rerun()

        if col_s2.button("🔲 Бүгдийг болиулах"):
            for r in st.session_state.records:
                r["Selected"] = False
            st.rerun()

        df_valid = pd.DataFrame(st.session_state.records)
        cols_to_show = ["Selected", "SeqNo", "PNLNo", "Passenger Name", "PNR", "Class", "StatusCode", "BookingDate", "OfficeCode", "TicketNo", "Email", "Language", "SendStatus"]
        
        edited_df = st.data_editor(
            df_valid[cols_to_show],
            column_config={
                "Selected": st.column_config.CheckboxColumn("Сонгох", default=True),
                "Language": st.column_config.SelectboxColumn(
                    "Хэл (Language)",
                    options=["MN", "EN"],
                    required=True
                )
            },
            disabled=["SeqNo", "PNLNo", "Passenger Name", "PNR", "Class", "StatusCode", "BookingDate", "OfficeCode", "TicketNo", "Email", "SendStatus"],
            hide_index=True,
            use_container_width=True,
            key="passenger_editor"
        )
        
        for idx, row in edited_df.iterrows():
            st.session_state.records[idx]["Selected"] = row["Selected"]
            st.session_state.records[idx]["Language"] = row["Language"]
    else:
        st.info("Одоогоор уншигдсан идэвхтэй зорчигч байхгүй байна.")

with tab2:
    if st.session_state.missing_records:
        df_missing = pd.DataFrame(st.session_state.missing_records)
        cols_missing = ["SeqNo", "PNLNo", "Passenger Name", "PNR", "Class", "StatusCode", "BookingDate", "OfficeCode", "TicketNo", "Email", "MissingReason"]
        st.dataframe(df_missing[cols_missing], hide_index=True, use_container_width=True)
    else:
        st.info("Дутуу мэдээлэлтэй зорчигч байхгүй байна.")

active_cnt = len(st.session_state.records)
missing_cnt = len(st.session_state.missing_records)
total_pax = active_cnt + missing_cnt

selected_passengers = [r for r in st.session_state.records if r.get("Selected", False)]
selected_cnt = len(selected_passengers)

mn_cnt = sum(1 for r in selected_passengers if r["Language"] == "MN")
en_cnt = sum(1 for r in selected_passengers if r["Language"] == "EN")

st.write(f"**Мэдээлэл:** 🔵 MN: {mn_cnt} | 🟢 EN: {en_cnt} | 🎯 Идэвхтэй: {selected_cnt}/{active_cnt} | ⚠️ Мэдээлэл Дутуу: {missing_cnt} | 👥 **Нийт зорчигчид:** {total_pax}")

with tab3:
    if selected_cnt == 0:
        st.warning("⚠️ Сонгосон зорчигч байхгүй байна. Та 'Идэвхтэй Зорчигчид' жагсаалтаас доод тал нь 1 зорчигч сонгоно уу.")
    else:
        col_p1, col_p2, col_p3 = st.columns([2, 1, 1])
        with col_p1:
            preview_lang = st.radio("Preview Хэл:", ["MN", "EN"], horizontal=True)
        with col_p2:
            st.write("")
            if st.button("✏ Засах" if not st.session_state.edit_mode else "👁 Харж шалгах"):
                st.session_state.edit_mode = not st.session_state.edit_mode
                st.rerun()
        with col_p3:
            st.write("")
            if st.session_state.custom_templates[preview_lang]["text"]:
                if st.button("🔄 Анхны хувилбар"):
                    st.session_state.custom_templates[preview_lang] = {"subject": "", "text": ""}
                    st.session_state.edit_mode = False
                    st.rerun()

        orig_subj, orig_text = generate_email_text_base(preview_lang, flight_info)

        curr_subj = st.session_state.custom_templates[preview_lang]["subject"] or orig_subj
        curr_text = st.session_state.custom_templates[preview_lang]["text"] or orig_text

        lbl_title = "ГАРЧИГ:" if preview_lang == "MN" else "SUBJECT:"

        if st.session_state.edit_mode:
            st.info("💡 Текст доторх `{PAX_NAME}`, `{PNR}`, `{TICKET_NO}` түлхүүр үгс нь зорчигч бүрийн мэдээллээр автоматаар солигдох болно.")
            new_subj = st.text_input(f"{lbl_title}", value=curr_subj)
            new_text = st.text_area("Засах боломжтой эх текст:", value=curr_text, height=320)
            
            st.session_state.custom_templates[preview_lang]["subject"] = new_subj
            st.session_state.custom_templates[preview_lang]["text"] = new_text
        else:
            sample_rec = selected_passengers[0] if selected_cnt >= 1 else None
            st.markdown(f"**{lbl_title}** {curr_subj}")
            
            rendered_plain = render_custom_template(curr_text, sample_rec)
            rendered_html = text_to_html(rendered_plain)
            st.components.v1.html(rendered_html, height=380, scrolling=True)

# TAB 4: ADVANCED AUDIT LOG HISTORY (Шүүлтүүр, Статистик, Экспорт)
with tab4:
    st.subheader("📜 Үйл ажиллагааны түүх ба Шүүлтүүр")
    
    if os.path.exists(LOGS_FILE):
        try:
            df_logs = pd.read_csv(LOGS_FILE, encoding="utf-8-sig", on_bad_lines="skip")
            
            if not df_logs.empty:
                col_m1, col_m2, col_m3, col_m4 = st.columns(4)
                col_m1.metric("Нийт үйлдэл", len(df_logs))
                col_m2.metric("Нэвтэрсэн", len(df_logs[df_logs["Action"] == "Нэвтэрсэн"]))
                col_m3.metric("Имэйл илгээсэн", len(df_logs[df_logs["Action"] == "Имэйл илгээсэн"]))
                col_m4.metric("Excel татсан", len(df_logs[df_logs["Action"] == "Excel тайлан татсан"]))
                
                st.divider()
                
                f_c1, f_c2 = st.columns(2)
                all_users_list = ["Бүгд"] + sorted(df_logs["Username"].unique().tolist()) if "Username" in df_logs.columns else ["Бүгд"]
                all_actions_list = ["Бүгд"] + sorted(df_logs["Action"].unique().tolist()) if "Action" in df_logs.columns else ["Бүгд"]
                
                selected_user_filter = f_c1.selectbox("Хэрэглэгчээр шүүх:", all_users_list)
                selected_action_filter = f_c2.selectbox("Үйлдлийн төрлөөр шүүх:", all_actions_list)
                
                filtered_df = df_logs.copy()
                if selected_user_filter != "Бүгд":
                    filtered_df = filtered_df[filtered_df["Username"] == selected_user_filter]
                if selected_action_filter != "Бүгд":
                    filtered_df = filtered_df[filtered_df["Action"] == selected_action_filter]
                
                st.dataframe(filtered_df.sort_index(ascending=False), use_container_width=True)
                
                log_excel_buffer = io.BytesIO()
                filtered_df.to_excel(log_excel_buffer, index=False, engine="openpyxl")
                log_excel_buffer.seek(0)
                
                st.download_button(
                    label="📥 Үйл ажиллагааны түүхийг Excel болгон татах",
                    data=log_excel_buffer,
                    file_name=f"Audit_Logs_{datetime.now(ZoneInfo('Asia/Ulaanbaatar')).strftime('%Y%m%d_%H%M%S')}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
            else:
                st.info("Лог түүх хоосон байна.")
        except Exception:
            st.warning("Лог файлыг уншихад алдаа гарлаа. `logs.csv` файлыг шинээр үүсгэнэ үү.")
    else:
        st.info("Одоогоор лог түүх үүсээгүй байна.")

# TAB 5: USER MANAGEMENT
with tab5:
    st.subheader("👥 Хэрэглэгчийн бүртгэл ба тохиргоо")
    users = load_users()
    
    if st.session_state.user_info["role"] == "admin":
        st.markdown("#### ➕ Шинэ хэрэглэгч нэмэх")
        
        with st.form("add_user_form", clear_on_submit=True):
            c_u1, c_u2, c_u3 = st.columns([2, 2, 1])
            new_uname = c_u1.text_input("Нэвтрэх нэр", key="form_new_uname")
            new_pass = c_u2.text_input("Нууц үг", type="password", key="form_new_pass")
            new_role = c_u3.selectbox("Эрх", ["user", "admin"], key="form_new_role")
            
            submitted = st.form_submit_button("Хэрэглэгч нэмэх", type="primary")
            if submitted:
                if new_uname and new_pass:
                    if new_uname in users:
                        st.error("Ийм нэвтрэх нэртэй хэрэглэгч аль хэдийн байна!")
                    else:
                        users[new_uname] = {
                            "password_hash": hash_password(new_pass),
                            "role": new_role
                        }
                        save_users(users)
                        add_log(st.session_state.user_info['username'], "Хэрэглэгч нэмсэн", details=f"Хэрэглэгч '{new_uname}' нэмэгдлээ")
                        st.success(f"Хэрэглэгч '{new_uname}' амжилттай нэмэгдлээ!")
                        time.sleep(1)
                        st.rerun()
                else:
                    st.warning("Нэвтрэх нэр болон нууц үгээ оруулна уу.")
        
        st.divider()
        st.markdown("#### 📋 Бүртгэлтэй хэрэглэгчдийн жагсаалт ба удирдлага")
        
        for uname, udata in list(users.items()):
            with st.expander(f"👤 {uname} (Эрх: {udata['role']})"):
                u_col1, u_col2 = st.columns([3, 1])
                
                with u_col1:
                    new_reset_pass = st.text_input(f"Шинэ нууц үг тохируулах ({uname}):", type="password", key=f"reset_pass_{uname}")
                    if st.button(f"🔑 Нууц үг солих", key=f"btn_reset_{uname}"):
                        if new_reset_pass:
                            users[uname]["password_hash"] = hash_password(new_reset_pass)
                            save_users(users)
                            add_log(st.session_state.user_info['username'], "Нууц үг шинэчилсэн", details=f"Хэрэглэгч '{uname}'-ийн нууц үгийг солов.")
                            
                            st.success(f"✅ '{uname}' хэрэглэгчийн нууц үг амжилттай солигдлоо!")
                            time.sleep(1.5)
                            st.rerun()
                        else:
                            st.warning("⚠️ Шинэ нууц үгээ оруулна уу.")

                with u_col2:
                    st.write("")
                    st.write("")
                    if uname == st.session_state.user_info['username']:
                        st.caption("(Одоо нэвтэрсэн)")
                    else:
                        if st.button("🗑️ Устгах", key=f"del_{uname}"):
                            del users[uname]
                            save_users(users)
                            add_log(st.session_state.user_info['username'], "Хэрэглэгч устгасан", details=f"Хэрэглэгч '{uname}' устгагдлаа")
                            st.success(f"Хэрэглэгч '{uname}' устгагдлаа!")
                            time.sleep(1)
                            st.rerun()
    else:
        st.info("Хэрэглэгч нэмэх, нууц үг шинэчлэх болон устгах эрх зөвхөн АДМИН хэрэглэгчид боломжтой.")

st.divider()

# --- DIALOG / MODAL FOR CONFIRMATION ---
@st.dialog("Имэйл текстийг шалгах ба Баталгаажуулах", width="large")
def confirm_and_send_dialog():
    st.warning("⚠ Дараах имэйлийн эх текст зорчигчид руу илгээгдэх гэж байна. Шалгаад 'Илгээх' эсвэл 'Засах' товчийг сонгоно уу.")
    
    tab_mn, tab_en = st.tabs(["🇲🇳 Монгол (MN)", "🇬🇧 Англи (EN)"])
    sample_rec = selected_passengers[0] if len(selected_passengers) >= 1 else None
    
    with tab_mn:
        orig_subj, orig_text = generate_email_text_base("MN", flight_info)
        s_subj = st.session_state.custom_templates["MN"]["subject"] or orig_subj
        s_text = st.session_state.custom_templates["MN"]["text"] or orig_text
        st.markdown(f"**ГАРЧИГ:** {s_subj}")
        st.components.v1.html(text_to_html(render_custom_template(s_text, sample_rec)), height=280, scrolling=True)
        
    with tab_en:
        orig_subj_en, orig_text_en = generate_email_text_base("EN", flight_info)
        s_subj_en = st.session_state.custom_templates["EN"]["subject"] or orig_subj_en
        s_text_en = st.session_state.custom_templates["EN"]["text"] or orig_text_en
        st.markdown(f"**SUBJECT:** {s_subj_en}")
        st.components.v1.html(text_to_html(render_custom_template(s_text_en, sample_rec)), height=280, scrolling=True)
        
    col_d1, col_d2 = st.columns([1, 1])
    if col_d1.button("✅ Зөв, одоо илгээх", type="primary", use_container_width=True):
        st.session_state.start_send_process = True
        st.rerun()
        
    if col_d2.button("✏️ Засах шаардлагатай", use_container_width=True):
        st.rerun()

# --- ACTIONS: SEND & EXPORT ---
col_act1, col_act2 = st.columns([2, 1])

with col_act1:
    if st.button("🚀 СОНГОСОН ЗОРЧИГЧИДОД ИМЭЙЛ ИЛГЭЭХ", type="primary", use_container_width=True):
        if not sender_email or not app_password:
            st.error("Систем нэвтрэх Gmail хаяг болон App Password оруулаагүй байна! Админ хэрэглэгчээр тохиргоог хадгална уу.")
        elif not st.session_state.records:
            st.warning("Илгээх зорчигч байхгүй байна.")
        elif selected_cnt == 0:
            st.warning("Нэг ч зорчигч сонгогдоогүй байна.")
        else:
            confirm_and_send_dialog()

    if st.session_state.get("start_send_process", False):
        st.session_state.start_send_process = False
        success_count, fail_count = 0, 0
        
        progress_bar = st.progress(0)
        status_text = st.empty()
        
        for i, pax in enumerate(selected_passengers):
            lang = pax["Language"]
            pax_name = pax.get("Passenger Name", "Passenger")

            for recipient in pax.get("EmailList", []):
                try:
                    orig_subj, orig_text = generate_email_text_base(lang, flight_info)
                    
                    cust_subj = st.session_state.custom_templates[lang]["subject"]
                    cust_text = st.session_state.custom_templates[lang]["text"]
                    
                    final_subj = cust_subj if cust_subj else orig_subj
                    raw_text = cust_text if cust_text else orig_text
                    
                    final_plain = render_custom_template(raw_text, pax)
                    final_html = text_to_html(final_plain)

                    send_email_smtp(
                        sender_email, 
                        app_password, 
                        recipient, 
                        pax_name, 
                        final_subj, 
                        final_plain, 
                        final_html
                    )
                    
                    pax["SendStatus"] = "Sent Successfully"
                    pax["SentTime"] = get_ubn_now()
                    success_count += 1
                    
                    sleep_time = random.uniform(2.0, 4.5)
                    status_text.text(f"Илгээж байна ({i+1}/{len(selected_passengers)}): {recipient} ... ({sleep_time:.1f}с хүлээж байна)")
                    time.sleep(sleep_time)

                    if (i + 1) % 15 == 0 and i + 1 < len(selected_passengers):
                        status_text.text(f"⏳ Серверийн ачааллыг багасгахад 12 секунд хүлээж байна...")
                        time.sleep(12)

                except Exception as e:
                    pax["SendStatus"] = f"Failed: {str(e)}"
                    fail_count += 1
            
            progress_bar.progress((i + 1) / len(selected_passengers))

        status_text.empty()
        
        add_log(
            st.session_state.user_info['username'], 
            "Имэйл илгээсэн", 
            flight_no=flt_no, 
            flight_date=flt_date,
            details=f"Нислэгийн огноо: {flt_date}, Амжилттай илгээсэн: {success_count} зорчигч, Алдаа гарсан: {fail_count}"
        )
        
        st.success(f"Ажиллагаа дууслаа! Нийт амжилттай: {success_count}, Амжилтгүй: {fail_count}")
        st.rerun()

with col_act2:
    all_data = st.session_state.records + st.session_state.missing_records
    if all_data:
        flt_val = st.session_state.input_flt_no.strip()
        date_val = st.session_state.input_flt_date.strip()
        
        if flt_val and date_val:
            excel_filename = f"{flt_val}_{date_val}.xlsx"
        elif flt_val:
            excel_filename = f"{flt_val}_Report.xlsx"
        else:
            ubn_file_time = datetime.now(ZoneInfo("Asia/Ulaanbaatar")).strftime("%Y%m%d_%H%M%S")
            excel_filename = f"PNL_Report_{ubn_file_time}.xlsx"

        excel_data = create_formatted_excel(st.session_state.records, st.session_state.missing_records)

        if st.download_button(
            label="📥 EXCEL тайлан татах",
            data=excel_data,
            file_name=excel_filename,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True
        ):
            add_log(
                st.session_state.user_info['username'],
                "Excel тайлан татсан",
                flight_no=flt_val if flt_val else "-",
                flight_date=date_val if date_val else "-",
                details=f"Нийт зорчигчдын тайлан татсан (Файлын нэр: {excel_filename})"
            )
